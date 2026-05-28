from __future__ import annotations

import csv
import re
from dataclasses import dataclass, asdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from .config import settings

SUPPORTED_SUFFIXES = {".pdf", ".csv", ".xlsx"}


@dataclass
class Chunk:
    text: str
    source: str          # file name
    locator: str         # e.g. "صفحة 12" or "صف 4"
    chunk_id: int

    def to_dict(self) -> dict:
        return asdict(self)


_WS = re.compile(r"[ \t ]+")
_NL = re.compile(r"\n{3,}")


def _clean(text: str) -> str:
    text = text.replace("\r", "\n")
    text = _WS.sub(" ", text)
    text = _NL.sub("\n\n", text)
    return text.strip()


def _split_text(text: str, size: int, overlap: int) -> list[str]:
    """Split on paragraph boundaries, packing into ~size chunks with overlap."""
    text = _clean(text)
    if not text:
        return []

    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        if len(buf) + len(para) + 1 <= size:
            buf = f"{buf}\n{para}" if buf else para
            continue
        if buf:
            chunks.append(buf)
        # If a single paragraph is larger than size, hard-split it.
        if len(para) > size:
            for i in range(0, len(para), size - overlap):
                chunks.append(para[i : i + size])
            buf = ""
        else:
            buf = para
    if buf:
        chunks.append(buf)

    # Apply overlap between adjacent chunks for better retrieval continuity.
    if overlap > 0 and len(chunks) > 1:
        overlapped: list[str] = [chunks[0]]
        for prev, cur in zip(chunks, chunks[1:]):
            tail = prev[-overlap:]
            overlapped.append(f"{tail} {cur}".strip())
        chunks = overlapped

    return [c for c in chunks if c.strip()]


def _read_pdf(path: Path) -> list[tuple[str, str]]:
    """Return list of (page_text, locator)."""
    reader = PdfReader(str(path))
    pages: list[tuple[str, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append((text, f"صفحة {i}"))
    return pages


def _read_csv(path: Path) -> list[tuple[str, str]]:
    """Each row becomes a 'col: value' text block. Returns (row_text, locator)."""
    rows: list[tuple[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader, start=1):
            parts = [f"{k}: {v}" for k, v in row.items() if v and str(v).strip()]
            if parts:
                rows.append(("\n".join(parts), f"صف {i}"))
    return rows


def _read_xlsx(path: Path) -> list[tuple[str, str]]:
    """Each row of every sheet becomes a 'col: value' text block."""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    rows: list[tuple[str, str]] = []
    for ws in wb.worksheets:
        iterator = ws.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            continue
        headers = [str(h).strip() if h is not None else f"عمود{i+1}" for i, h in enumerate(header)]
        sheet_tag = f"{ws.title}!" if len(wb.worksheets) > 1 else ""
        for i, row in enumerate(iterator, start=2):
            parts = [
                f"{headers[j]}: {val}"
                for j, val in enumerate(row)
                if j < len(headers) and val is not None and str(val).strip()
            ]
            if parts:
                rows.append(("\n".join(parts), f"{sheet_tag}صف {i}"))
    wb.close()
    return rows


def load_chunks(data_dir: Path | None = None) -> list[Chunk]:
    data_dir = data_dir or settings.data_dir
    chunks: list[Chunk] = []
    cid = 0

    files = sorted(
        p for p in data_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )

    for path in files:
        name = path.name
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            segments = _read_pdf(path)
        elif suffix == ".xlsx":
            segments = _read_xlsx(path)
        else:
            segments = _read_csv(path)

        for seg_text, locator in segments:
            for piece in _split_text(seg_text, settings.chunk_size, settings.chunk_overlap):
                chunks.append(Chunk(text=piece, source=name, locator=locator, chunk_id=cid))
                cid += 1

    return chunks
